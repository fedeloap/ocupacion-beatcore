"""Excel con Lagree y Pilates Reformer en Chamberí y Salamanca, sacado de los datos de la web (public/data).
Uso:  python informe_excel.py 2026-10-11 2026-10-16
"""
import collections
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DESDE = sys.argv[1] if len(sys.argv) > 1 else "2026-10-11"
HASTA = sys.argv[2] if len(sys.argv) > 2 else "2026-10-16"
DISCS = [("Lagree", "Lagree (incluye BeatCore)"), ("Reformer", "Pilates Reformer")]
ZONAS = ["Chamberí", "Salamanca"]

idx = json.load(open("public/data/clases.json", encoding="utf-8"))
est, dic = idx["estudios"], idx["dic"]
acc = collections.defaultdict(lambda: [0, 0, 0.0, 0.0])      # clases, con reservas, reservas, plazas
for m in idx["meses"]:
    d = json.load(open(f"public/data/clases_{m['mes']}.json", encoding="utf-8"))
    for j in range(d["n"]):
        f = dic["fecha"][d["f"][j]]
        if not (DESDE <= f <= HASTA):
            continue
        z = dic["centro"][d["c"][j]]
        disc = dic["disc"][d["d"][j]]
        if z not in ZONAS or disc not in ("Lagree", "Reformer"):
            continue
        e = est[d["e"][j]]
        x = acc[(disc, z, e["n"], e["p"])]
        x[0] += 1
        if d["r"][j] is not None and d["a"][j]:
            x[1] += 1
            x[2] += d["r"][j]
            x[3] += d["a"][j]

wb = Workbook()
AZUL = PatternFill("solid", fgColor="1F4E78")
GRIS = PatternFill("solid", fgColor="E7E6E6")
VERDE = PatternFill("solid", fgColor="E2EFDA")
NEG = Font(bold=True)
BLANCO = Font(bold=True, color="FFFFFF")


def cabecera(ws, fila, textos):
    for i, t in enumerate(textos, 1):
        c = ws.cell(fila, i, t)
        c.font, c.fill = BLANCO, AZUL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def ajustar(ws, anchos):
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = a


def media(x):
    return x[2] / x[1] if x[1] else None


def plazas(x):
    return x[3] / x[1] if x[1] else None


def ocup(x):
    return x[2] / x[3] if x[3] else None


def sumar(filas):
    t = [0, 0, 0.0, 0.0]
    for f in filas:
        for k in range(4):
            t[k] += f[k]
    return t


def fila_datos(ws, r, nombre, x, base_disc=None, base_zona=None, negrita=False, relleno=None):
    vals = [nombre, x[0], x[2] if x[1] else None, media(x), plazas(x), ocup(x),
            (x[2] / base_disc) if (base_disc and x[1]) else None, (x[2] / base_zona) if (base_zona and x[1]) else None]
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v)
        if negrita:
            c.font = NEG
        if relleno:
            c.fill = relleno
    for col, fmt in ((4, "0.0"), (5, "0.0"), (6, "0%"), (7, "0.0%"), (8, "0.0%")):
        ws.cell(r, col).number_format = fmt


COLS = ["Centro", "Clases dadas", "Reservas", "Media de gente por clase", "Plazas por clase", "Ocupación",
        "% de las reservas de la disciplina en la zona", "% de las reservas de las dos disciplinas en la zona"]

# ---------------- Resumen
ws = wb.active
ws.title = "Resumen"
ws["A1"] = "Lagree y Pilates Reformer en Chamberí y Salamanca"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = f"Periodo: del {DESDE} al {HASTA}. Datos de la web (clases ya dadas). BeatCore está dentro de Lagree. Reservas = gente apuntada a clases ya dadas, no asistencia."
ws["A3"] = "Media por clase, plazas por clase y ocupación se calculan sobre las clases con reservas conocidas (los centros de bsport no publican reservas)."
r = 5
ws.cell(r, 1, "Total de las dos zonas").font = NEG
r += 1
cabecera(ws, r, ["Disciplina", "Clases dadas", "Reservas", "Media de gente por clase", "Plazas por clase", "Ocupación", "% de las reservas de las dos disciplinas"])
tot = {d: sumar([v for k, v in acc.items() if k[0] == d]) for d, _ in DISCS}
gran = sum(t[2] for t in tot.values())
for d, nom in DISCS:
    r += 1
    t = tot[d]
    for i, v in enumerate([nom, t[0], t[2], media(t), plazas(t), ocup(t), t[2] / gran if gran else None], 1):
        ws.cell(r, i, v)
    for col, fmt in ((4, "0.0"), (5, "0.0"), (6, "0%"), (7, "0.0%")):
        ws.cell(r, col).number_format = fmt
r += 1
t = sumar(list(tot.values()))
for i, v in enumerate(["Total", t[0], t[2], media(t), plazas(t), ocup(t), 1.0], 1):
    c = ws.cell(r, i, v)
    c.font = NEG
    c.fill = GRIS
for col, fmt in ((4, "0.0"), (5, "0.0"), (6, "0%"), (7, "0.0%")):
    ws.cell(r, col).number_format = fmt

r += 2
ws.cell(r, 1, "Total de cada zona").font = NEG
r += 1
cabecera(ws, r, ["Zona", "Disciplina", "Clases dadas", "Reservas", "Media de gente por clase", "Plazas por clase", "Ocupación",
                 "% de la disciplina (dos zonas)", "% de la zona (dos disciplinas)"])
for z in ZONAS:
    zt = sumar([v for k, v in acc.items() if k[1] == z])
    for d, nom in DISCS:
        r += 1
        x = sumar([v for k, v in acc.items() if k[0] == d and k[1] == z])
        vals = [z, nom, x[0], x[2], media(x), plazas(x), ocup(x), x[2] / tot[d][2] if tot[d][2] else None, x[2] / zt[2] if zt[2] else None]
        for i, v in enumerate(vals, 1):
            ws.cell(r, i, v)
        for col, fmt in ((5, "0.0"), (6, "0.0"), (7, "0%"), (8, "0.0%"), (9, "0.0%")):
            ws.cell(r, col).number_format = fmt
    r += 1
    vals = [z, "Total zona", zt[0], zt[2], media(zt), plazas(zt), ocup(zt), None, zt[2] / gran if gran else None]
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v)
        c.font, c.fill = NEG, GRIS
    for col, fmt in ((5, "0.0"), (6, "0.0"), (7, "0%"), (9, "0.0%")):
        ws.cell(r, col).number_format = fmt
ajustar(ws, [26, 24, 14, 14, 18, 16, 14, 24, 26])
ws.row_dimensions[6].height = 45

# ---------------- una hoja por zona
for z in ZONAS:
    ws = wb.create_sheet(z)
    ws["A1"] = f"{z}: centros por disciplina (ordenados por reservas)"
    ws["A1"].font = Font(bold=True, size=13)
    zt = sumar([v for k, v in acc.items() if k[1] == z])
    r = 3
    for d, nom in DISCS:
        filas = sorted([(k[2], k[3], v) for k, v in acc.items() if k[0] == d and k[1] == z], key=lambda t: -t[2][2])
        dt = sumar([v for _, _, v in filas])
        ws.cell(r, 1, nom).font = Font(bold=True, size=12)
        r += 1
        cabecera(ws, r, COLS)
        ws.row_dimensions[r].height = 55
        for n, (nombre, propio, v) in enumerate(filas):
            r += 1
            fila_datos(ws, r, nombre + (" (BeatCore)" if propio else ""), v, dt[2], zt[2],
                       relleno=VERDE if n < 5 and v[1] else None)
        r += 1
        fila_datos(ws, r, f"Total {nom} en {z}", dt, dt[2], zt[2], negrita=True, relleno=GRIS)
        r += 2
    ws.cell(r, 1, "Verde = los 5 centros con más reservas de esa disciplina en la zona.")
    ajustar(ws, [44, 14, 12, 16, 14, 12, 22, 24])

# ---------------- Top 5
ws = wb.create_sheet("Top 5")
ws["A1"] = "Los 5 centros con más reservas, por zona y disciplina"
ws["A1"].font = Font(bold=True, size=13)
r = 3
for z in ZONAS:
    zt = sumar([v for k, v in acc.items() if k[1] == z])
    for d, nom in DISCS:
        filas = sorted([(k[2], k[3], v) for k, v in acc.items() if k[0] == d and k[1] == z], key=lambda t: -t[2][2])
        dt = sumar([v for _, _, v in filas])
        ws.cell(r, 1, f"{z} · {nom}").font = Font(bold=True, size=12)
        r += 1
        cabecera(ws, r, ["Puesto"] + COLS)
        ws.row_dimensions[r].height = 55
        for n, (nombre, propio, v) in enumerate(filas[:5], 1):
            r += 1
            ws.cell(r, 1, n)
            vals = [nombre + (" (BeatCore)" if propio else ""), v[0], v[2] if v[1] else None, media(v), plazas(v), ocup(v),
                    v[2] / dt[2] if dt[2] and v[1] else None, v[2] / zt[2] if zt[2] and v[1] else None]
            for i, x in enumerate(vals, 2):
                ws.cell(r, i, x)
            for col, fmt in ((5, "0.0"), (6, "0.0"), (7, "0%"), (8, "0.0%"), (9, "0.0%")):
                ws.cell(r, col).number_format = fmt
        r += 1
        resto = filas[5:]
        if resto:
            rt = sumar([v for _, _, v in resto])
            ws.cell(r, 2, f"Resto ({len(resto)} centros)")
            ws.cell(r, 3, rt[0])
            ws.cell(r, 4, rt[2])
            ws.cell(r, 8, rt[2] / dt[2] if dt[2] else None).number_format = "0.0%"
            ws.cell(r, 9, rt[2] / zt[2] if zt[2] else None).number_format = "0.0%"
            r += 1
        r += 1
ajustar(ws, [8, 44, 14, 12, 16, 14, 12, 22, 24])

os.makedirs("informes", exist_ok=True)
ruta = os.path.join("informes", f"lagree_reformer_chamberi_salamanca_{DESDE}_{HASTA}.xlsx")
wb.save(ruta)
print(ruta)
